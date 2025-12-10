import json
from datetime import timedelta

import jdatetime
import openpyxl
from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import format_html
from jalali_date import datetime2jalali
from jalali_date.admin import ModelAdminJalaliMixin
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .models import (
    UserActivity, UserStatistics, ReservationLog, PaymentLog, ErrorLog
)


def datetime2jalali(date_time):
    if not date_time:
        return '-'
    if isinstance(date_time, type(jdatetime.date.today())):
        jdate = jdatetime.date.fromgregorian(date=date_time)
        return jdate.strftime('%Y/%m/%d')

    jdate = jdatetime.datetime.fromgregorian(datetime=date_time)
    return jdate.strftime('%Y/%m/%d - %H:%M')


def date2jalali(date_obj):
    if not date_obj:
        return '-'
    jdate = jdatetime.date.fromgregorian(date=date_obj)
    return jdate.strftime('%Y/%m/%d')


@admin.register(UserActivity)
class UserActivityAdmin(ModelAdminJalaliMixin, admin.ModelAdmin):
    """
    Comprehensive user activity admin
    """

    list_display = [
        'id',
        'user_display',
        'action_badge',
        'severity_badge',
        'description_short',
        'device_info',
        'ip_display',
        'get_jalali_created_at'
    ]

    list_filter = [
        'action_type',
        'severity',
        'device_type',
        'browser',
        'os',
        'created_at'
    ]

    search_fields = [
        'user__full_name',
        'user__phone',
        'description',
        'ip_address',
        'session_key'
    ]

    readonly_fields = [
        'user',
        'session_key',
        'action_type',
        'severity',
        'description',
        'content_type',
        'object_id',
        'metadata_display',
        'ip_address',
        'user_agent',
        'device_type',
        'browser',
        'os',
        'created_at',
        'get_jalali_created_at'
    ]

    date_hierarchy = 'created_at'
    list_per_page = 50

    fieldsets = (
        ('اطلاعات کاربر', {
            'fields': ('user', 'session_key')
        }),
        ('جزئیات عملیات', {
            'fields': ('action_type', 'severity', 'description', 'content_type', 'object_id')
        }),
        ('اطلاعات درخواست', {
            'fields': ('ip_address', 'user_agent', 'device_type', 'browser', 'os'),
            'classes': ('collapse',)
        }),
        ('داده‌های اضافی', {
            'fields': ('metadata_display',),
            'classes': ('collapse',)
        }),
        ('زمان', {
            'fields': ('created_at', 'get_jalali_created_at')
        })
    )
    actions = ['export_as_csv', 'export_as_excel']

    def export_as_excel(modeladmin, request, queryset):
        """
        Export UserActivity as styled Excel (.xlsx) with alternating row colors
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "فعالیت کاربران"

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3A75B8", end_color="3A75B8", fill_type="solid")  # darker sky blue
        right_alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Alternating row colors
        row_colors = ["FFFFFF", "D9EAF7"]  # white, sky blue

        # Headers
        headers = ['ID', 'کاربر', 'نوع عملیات', 'توضیحات', 'دستگاه', 'IP', 'زمان']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = right_alignment
            cell.border = thin_border

        # Rows
        for row_num, obj in enumerate(queryset.order_by('created_at'), start=2):
            user_text = f"{obj.user.full_name} - {obj.user.phone}" if obj.user else (
                obj.session_key[:10] if obj.session_key else '-')
            device_text = f"{obj.device_type or '-'} - {obj.browser or '-'} - {obj.os or '-'}"

            row = [
                obj.id,
                user_text,
                obj.get_action_type_display(),
                obj.description,
                device_text,
                obj.ip_address or '-',
                modeladmin.get_jalali_created_at(obj)
            ]

            # Alternate row color
            fill_color = PatternFill(start_color=row_colors[(row_num - 2) % 2], end_color=row_colors[(row_num - 2) % 2],
                                     fill_type="solid")

            for col_num, value in enumerate(row, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.alignment = right_alignment
                cell.fill = fill_color
                cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 5  # padding

        # Response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="user_activities.xlsx"'
        wb.save(response)
        return response

    export_as_excel.short_description = "صدور Excel با استایل و ردیف‌های رنگی"

    def user_display(self, obj):
        """Display user with link"""
        if obj.user:
            url = f'/admin/accounts/user/{obj.user.id}/change/'
            return format_html(
                '<a href="{}" style="font-weight: bold;">{}</a><br/>'
                '<small style="color: #6b7280;">{}</small>',
                url,
                obj.user.full_name,
                obj.user.phone
            )
        return format_html(
            '<span style="color: #9ca3af;">Guest</span><br/>'
            '<small style="color: #d1d5db;">{}</small>',
            obj.session_key[:10] if obj.session_key else '-'
        )

    user_display.short_description = 'کاربر'

    def action_badge(self, obj):
        """Display action type with icon"""
        icons = {
            'registration': '📝',
            'login': '🔑',
            'logout': '🚪',
            'otp_request': '📱',
            'otp_verify': '✅',
            'reservation_start': '📅',
            'reservation_created': '✓',
            'reservation_cancelled': '❌',
            'payment_init': '💳',
            'payment_success': '💰',
            'payment_failed': '⚠️',
            'file_upload': '📎',
            'pdf_download': '📄',
        }
        icon = icons.get(obj.action_type, '📊')

        colors = {
            'login': '#10b981',
            'registration': '#3b82f6',
            'payment_success': '#059669',
            'payment_failed': '#ef4444',
            'reservation_cancelled': '#f59e0b',
        }
        color = colors.get(obj.action_type, '#6b7280')

        return format_html(
            '<span style="background: {}; color: white; padding: 4px 12px; '
            'border-radius: 8px; font-size: 11px; font-weight: 600;">'
            '{} {}</span>',
            color,
            icon,
            obj.get_action_type_display()
        )

    action_badge.short_description = 'عملیات'

    def severity_badge(self, obj):
        """Display severity level"""
        colors = {
            'info': '#3b82f6',
            'warning': '#f59e0b',
            'error': '#ef4444',
            'critical': '#dc2626'
        }
        icons = {
            'info': 'ℹ️',
            'warning': '⚠️',
            'error': '❌',
            'critical': '🔥'
        }

        return format_html(
            '<span style="background: {}; color: white; padding: 3px 10px; '
            'border-radius: 6px; font-size: 10px;">{} {}</span>',
            colors.get(obj.severity, '#6b7280'),
            icons.get(obj.severity, 'ℹ️'),
            obj.get_severity_display()
        )

    severity_badge.short_description = 'سطح'

    def description_short(self, obj):
        """Truncated description"""
        if len(obj.description) > 60:
            return obj.description[:60] + '...'
        return obj.description

    description_short.short_description = 'توضیحات'

    def device_info(self, obj):
        """Display device and browser info"""
        if obj.device_type or obj.browser:
            return format_html(
                '<strong>{}</strong><br/>'
                '<small style="color: #6b7280;">{} - {}</small>',
                obj.device_type or '-',
                obj.browser or '-',
                obj.os or '-'
            )
        return '-'

    device_info.short_description = 'دستگاه'

    def ip_display(self, obj):
        """Display IP address"""
        if obj.ip_address:
            return format_html(
                '<code style="background: #f3f4f6; padding: 3px 8px; '
                'border-radius: 4px; font-size: 11px;">{}</code>',
                obj.ip_address
            )
        return '-'

    ip_display.short_description = 'IP'

    def metadata_display(self, obj):
        """Display metadata as formatted JSON"""
        if obj.metadata:
            try:
                formatted = json.dumps(obj.metadata, indent=2, ensure_ascii=False)
                return format_html(
                    '<pre style="background: #f9fafb; padding: 10px; '
                    'border-radius: 5px; font-size: 12px; max-width: 600px; '
                    'overflow-x: auto;">{}</pre>',
                    formatted
                )
            except:
                return str(obj.metadata)
        return '-'

    metadata_display.short_description = 'داده‌های اضافی'

    @admin.display(description='تاریخ (شمسی)', ordering='created_at')
    def get_jalali_created_at(self, obj):
        return datetime2jalali(obj.created_at).strftime('%Y/%m/%d - %H:%M:%S')

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False


@admin.register(UserStatistics)
class UserStatisticsAdmin(ModelAdminJalaliMixin, admin.ModelAdmin):
    """
    User statistics dashboard
    """

    list_display = [
        'user_display',
        'reservation_stats',
        'payment_stats',
        'activity_stats',
        'last_activity_display'
    ]

    list_filter = [
        'last_activity',
        'last_reservation',
        'last_payment'
    ]

    search_fields = [
        'user__full_name',
        'user__phone',
        'user__email'
    ]

    readonly_fields = [
        'user',
        'total_reservations',
        'completed_reservations',
        'cancelled_reservations',
        'pending_reservations',
        'total_payments',
        'successful_payments',
        'failed_payments',
        'refunded_amount',
        'total_logins',
        'total_activities',
        'most_used_service',
        'first_activity',
        'last_activity',
        'last_reservation',
        'last_payment',
        'updated_at'
    ]

    fieldsets = (
        ('کاربر', {
            'fields': ('user',)
        }),
        ('آمار رزروها', {
            'fields': (
                'total_reservations',
                'completed_reservations',
                'cancelled_reservations',
                'pending_reservations'
            )
        }),
        ('آمار پرداخت', {
            'fields': (
                'total_payments',
                'successful_payments',
                'failed_payments',
                'refunded_amount'
            )
        }),
        ('آمار فعالیت', {
            'fields': (
                'total_logins',
                'total_activities',
                'most_used_service'
            )
        }),
        ('تاریخ‌ها', {
            'fields': (
                'first_activity',
                'last_activity',
                'last_reservation',
                'last_payment',
                'updated_at'
            ),
            'classes': ('collapse',)
        })
    )
    actions = ['export_statistics_excel']

    def export_statistics_excel(modeladmin, request, queryset):
        """
        Export UserStatistics as styled Excel (.xlsx)
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "آمار کاربران"

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3A75B8", end_color="3A75B8", fill_type="solid")  # dark sky blue
        right_alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Alternating row colors
        row_colors = ["FFFFFF", "D9EAF7"]  # white / sky blue

        # Headers
        headers = [
            'نام کاربر',
            'تلفن کاربر',
            'کل رزروها',
            'رزروهای موفق',
            'رزروهای لغو شده',
            'پرداخت کل (تومان)',
            'پرداخت‌های موفق',
            'پرداخت‌های ناموفق',
            'ورودها',
            'فعالیت‌ها',
            'آخرین فعالیت'
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = right_alignment
            cell.border = thin_border

        # Rows
        for row_num, obj in enumerate(queryset.order_by('user__full_name'), start=2):
            user_name = obj.user.full_name if obj.user else '-'
            user_phone = obj.user.phone if obj.user else '-'

            last_activity_text = '-'
            if obj.last_activity:
                last_activity_text = datetime2jalali(obj.last_activity)

            row = [
                user_name,
                user_phone,
                obj.total_reservations or 0,
                obj.completed_reservations or 0,
                obj.cancelled_reservations or 0,
                int(obj.total_payments or 0),
                obj.successful_payments or 0,
                obj.failed_payments or 0,
                obj.total_logins or 0,
                obj.total_activities or 0,
                last_activity_text
            ]

            fill_color = PatternFill(
                start_color=row_colors[(row_num - 2) % 2],
                end_color=row_colors[(row_num - 2) % 2],
                fill_type="solid"
            )

            for col_num, value in enumerate(row, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.alignment = right_alignment
                cell.fill = fill_color
                cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 5

        # Response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="user_statistics.xlsx"'
        wb.save(response)
        return response

    export_statistics_excel.short_description = "صدور Excel آمار کاربران با استایل"

    def user_display(self, obj):
        """Display user with link and avatar"""
        url = f'/admin/accounts/user/{obj.user.id}/change/'
        return format_html(
            '<div style="display: flex; align-items: center; gap: 10px;">'
            '<div style="width: 40px; height: 40px; background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); '
            'border-radius: 50%; display: flex; align-items: center; justify-content: center; '
            'color: white; font-weight: bold; font-size: 16px;">{}</div>'
            '<div>'
            '<a href="{}" style="font-weight: bold; font-size: 14px;">{}</a><br/>'
            '<small style="color: #6b7280;">{}</small>'
            '</div>'
            '</div>',
            obj.user.full_name[0].upper(),
            url,
            obj.user.full_name,
            obj.user.phone
        )

    user_display.short_description = 'کاربر'

    def reservation_stats(self, obj):
        """Display reservation statistics"""
        total = int(obj.total_reservations or 0)
        completed = int(obj.completed_reservations or 0)
        cancelled = int(obj.cancelled_reservations or 0)

        completed_pct = round((completed / total * 100) if total else 0)
        cancelled_pct = round((cancelled / total * 100) if total else 0)

        # همه مقادیر به str تبدیل می‌شوند
        return format_html(
            '<div style="min-width: 150px;">'
            '<strong style="font-size: 24px; color: #667eea;">{}</strong> '
            '<small style="color: #9ca3af;">کل</small><br/>'
            '<div style="margin-top: 5px;">'
            '<span style="color: #10b981;">✓ {} ({}%)</span> | '
            '<span style="color: #ef4444;">✗ {} ({}%)</span>'
            '</div>'
            '</div>',
            total,
            completed,
            completed_pct,
            cancelled,
            cancelled_pct
        )

    reservation_stats.short_description = 'رزروها'

    def payment_stats(self, obj):
        """Display payment statistics"""
        total = int(obj.total_payments or 0)
        successful = int(obj.successful_payments or 0)
        failed = int(obj.failed_payments or 0)

        # اصلاح: فرمت {:,} را به {} تغییر دهید و عدد را در آرگومان فرمت کنید
        return format_html(
            '<div style="min-width: 150px;">'
            '<strong style="font-size: 18px; color: #059669;">{}</strong> '  # اینجا {:,} حذف شد
            '<small style="color: #9ca3af;">تومان</small><br/>'
            '<div style="margin-top: 5px;">'
            '<span style="color: #10b981;">✓ {}</span> | '
            '<span style="color: #ef4444;">✗ {}</span>'
            '</div>'
            '</div>',
            f"{total:,}",  # اینجا عدد را فرمت کنید
            successful,
            failed
        )

    payment_stats.short_description = 'پرداخت‌ها'

    def activity_stats(self, obj):
        """Display activity statistics"""
        total_logins = int(obj.total_logins or 0)
        total_activities = int(obj.total_activities or 0)

        return format_html(
            '<div>'
            '<div><strong>{}</strong> ورود</div>'
            '<div><strong>{}</strong> فعالیت</div>'
            '</div>',
            total_logins,
            total_activities
        )

    activity_stats.short_description = 'فعالیت‌ها'

    def last_activity_display(self, obj):
        """Display last activity time"""
        if obj.last_activity:
            jalali = datetime2jalali(obj.last_activity)
            time_diff = timezone.now() - obj.last_activity

            if time_diff < timedelta(hours=1):
                time_ago = f'{int(time_diff.total_seconds() / 60)} دقیقه پیش'
                color = '#10b981'
            elif time_diff < timedelta(days=1):
                time_ago = f'{int(time_diff.total_seconds() / 3600)} ساعت پیش'
                color = '#f59e0b'
            else:
                time_ago = f'{time_diff.days} روز پیش'
                color = '#6b7280'

            return format_html(
                '<span style="color: {}; font-weight: 600;">{}</span><br/>'
                '<small style="color: #9ca3af;">{}</small>',
                color,
                time_ago,
                jalali
            )
        return '-'

    last_activity_display.short_description = 'آخرین فعالیت'

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False


@admin.register(ReservationLog)
class ReservationLogAdmin(ModelAdminJalaliMixin, admin.ModelAdmin):
    """
    Reservation changes tracking
    """
    list_display = [
        'id',
        'reservation_display',
        'user_display',
        'status_change',
        'changed_by_display',
        'get_jalali_created_at'
    ]

    list_filter = [
        'new_status',
        'created_at'
    ]

    search_fields = [
        'reservation__tracking_code',
        'user__full_name',
        'change_reason'
    ]

    readonly_fields = [
        'reservation',
        'user',
        'old_status',
        'new_status',
        'changed_by',
        'change_reason',
        'metadata_display',
        'created_at'
    ]

    date_hierarchy = 'created_at'

    fieldsets = (
        ('رزرو', {
            'fields': ('reservation', 'user')
        }),
        ('تغییر وضعیت', {
            'fields': ('old_status', 'new_status', 'changed_by', 'change_reason')
        }),
        ('اطلاعات اضافی', {
            'fields': ('metadata_display',),
            'classes': ('collapse',)
        }),
        ('زمان', {
            'fields': ('created_at',)
        })
    )

    def reservation_display(self, obj):
        """Display reservation with link"""
        url = f'/admin/council/reservation/{obj.reservation.id}/change/'
        return format_html(
            '<a href="{}" style="font-weight: bold;">{}</a><br/>'
            '<small style="color: #6b7280;">{}</small>',
            url,
            obj.reservation.tracking_code,
            obj.reservation.full_name
        )

    reservation_display.short_description = 'رزرو'

    def user_display(self, obj):
        """Display user"""
        if obj.user:
            url = f'/admin/accounts/user/{obj.user.id}/change/'
            return format_html('<a href="{}">{}</a>', url, obj.user.full_name)
        return '-'

    user_display.short_description = 'کاربر'

    def status_change(self, obj):
        """Display status change with arrows"""
        return format_html(
            '<div style="display: flex; align-items: center; gap: 10px;">'
            '<span style="background: #f3f4f6; padding: 5px 10px; border-radius: 6px;">{}</span>'
            '<span style="color: #667eea;">→</span>'
            '<span style="background: #667eea; color: white; padding: 5px 10px; border-radius: 6px;">{}</span>'
            '</div>',
            obj.old_status or 'جدید',
            obj.get_new_status_display()
        )

    status_change.short_description = 'تغییر وضعیت'

    def changed_by_display(self, obj):
        """Display who made the change"""
        if obj.changed_by:
            return obj.changed_by.full_name
        return 'سیستم'

    changed_by_display.short_description = 'تغییر توسط'

    def metadata_display(self, obj):
        """Display metadata"""
        if obj.metadata:
            try:
                formatted = json.dumps(obj.metadata, indent=2, ensure_ascii=False)
                return format_html('<pre>{}</pre>', formatted)
            except:
                return str(obj.metadata)
        return '-'

    metadata_display.short_description = 'داده‌ها'

    @admin.display(description='تاریخ', ordering='created_at')
    def get_jalali_created_at(self, obj):
        return datetime2jalali(obj.created_at).strftime('%Y/%m/%d - %H:%M')

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False


@admin.register(PaymentLog)
class PaymentLogAdmin(ModelAdminJalaliMixin, admin.ModelAdmin):
    """
    Payment transactions tracking
    """
    list_display = [
        'id',
        'reservation_display',
        'transaction_badge',
        'amount_display',
        'gateway_display',
        'status_display',
        'ref_id',
        'get_jalali_created_at'
    ]

    list_filter = [
        'transaction_type',
        'gateway_status',
        'created_at'
    ]

    search_fields = [
        'reservation__tracking_code',
        'user__full_name',
        'ref_id',
        'authority'
    ]

    readonly_fields = [
        'payment',
        'reservation',
        'user',
        'transaction_type',
        'amount',
        'gateway_status',
        'gateway_message',
        'authority',
        'ref_id',
        'request_data_display',
        'response_data_display',
        'ip_address',
        'created_at'
    ]

    date_hierarchy = 'created_at'
    list_per_page = 50

    fieldsets = (
        ('مشخصات تراکنش', {
            'fields': ('payment', 'reservation', 'user', 'transaction_type')
        }),
        # ('مبلغ و درگاه', {
        #     'fields': ('amount', 'gateway')
        # }),
        ('پاسخ درگاه', {
            'fields': ('gateway_status', 'gateway_message', 'authority', 'ref_id')
        }),
        ('داده‌های درخواست', {
            'fields': ('request_data_display',),
            'classes': ('collapse',)
        }),
        ('داده‌های پاسخ', {
            'fields': ('response_data_display',),
            'classes': ('collapse',)
        }),
        ('اطلاعات سیستمی', {
            'fields': ('ip_address', 'created_at')
        })
    )

    def reservation_display(self, obj):
        """Display reservation"""
        url = f'/admin/council/reservation/{obj.reservation.id}/change/'
        return format_html(
            '<a href="{}" style="font-weight: bold;">{}</a>',
            url,
            obj.reservation.tracking_code
        )

    reservation_display.short_description = 'رزرو'

    def transaction_badge(self, obj):
        """Display transaction type"""
        colors = {
            'init': '#3b82f6',
            'redirect': '#8b5cf6',
            'callback': '#ec4899',
            'verify': '#f59e0b',
            'success': '#10b981',
            'failed': '#ef4444',
            'refund_request': '#f97316',
            'refunded': '#6b7280'
        }

        return format_html(
            '<span style="background: {}; color: white; padding: 4px 12px; '
            'border-radius: 8px; font-size: 11px; font-weight: 600;">{}</span>',
            colors.get(obj.transaction_type, '#6b7280'),
            obj.get_transaction_type_display()
        )

    transaction_badge.short_description = 'نوع تراکنش'

    def amount_display(self, obj):
        """Display amount"""
        # اصلاح: رفع خطای ValueError با استفاده از f-string قبل از format_html
        amount = int(obj.amount or 0)
        return format_html(
            '<strong style="color: #059669; font-size: 14px;">{}</strong> '
            '<small style="color: #6b7280;">تومان</small>',
            f"{amount:,}"
        )

    amount_display.short_description = 'مبلغ'

    def gateway_display(self, obj):
        """Display gateway"""
        # اصلاح: استفاده از gateway_name به جای gateway
        if obj.gateway_name:
            return format_html(
                '<span style="background: #f3f4f6; padding: 3px 10px; '
                'border-radius: 6px; font-size: 11px;">{}</span>',
                obj.gateway_name
            )
        return '-'

    gateway_display.short_description = 'درگاه'

    def status_display(self, obj):
        """Display gateway status"""
        if obj.gateway_status:
            color = '#10b981' if obj.gateway_status in ['OK', 'success', '100'] else '#ef4444'
            return format_html(
                '<span style="color: {}; font-weight: 600;">{}</span>',
                color,
                obj.gateway_status
            )
        return '-'

    status_display.short_description = 'وضعیت'

    def request_data_display(self, obj):
        """Display request data"""
        if obj.request_data:
            try:
                formatted = json.dumps(obj.request_data, indent=2, ensure_ascii=False)
                return format_html('<pre style="max-height: 400px; overflow-y: auto;">{}</pre>', formatted)
            except:
                return str(obj.request_data)
        return '-'

    request_data_display.short_description = 'داده‌های درخواست'

    def response_data_display(self, obj):
        """Display response data"""
        if obj.response_data:
            try:
                formatted = json.dumps(obj.response_data, indent=2, ensure_ascii=False)
                return format_html('<pre style="max-height: 400px; overflow-y: auto;">{}</pre>', formatted)
            except:
                return str(obj.response_data)
        return '-'

    response_data_display.short_description = 'داده‌های پاسخ'

    @admin.display(description='تاریخ', ordering='created_at')
    def get_jalali_created_at(self, obj):
        return datetime2jalali(obj.created_at).strftime('%Y/%m/%d - %H:%M:%S')

    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False


@admin.register(ErrorLog)
class ErrorLogAdmin(ModelAdminJalaliMixin, admin.ModelAdmin):
    """
    System errors tracking
    """
    list_display = [
        'id',
        'error_badge',
        'error_message_short',
        'user_display',
        'view_name',
        'resolution_status',
        'get_jalali_created_at'
    ]

    list_filter = [
        'error_type',
        'is_resolved',
        'created_at'
    ]

    search_fields = [
        'error_message',
        'view_name',
        'url',
        'user__full_name'
    ]

    readonly_fields = [
        'error_type',
        'error_message',
        'stack_trace_display',
        'user',
        'url',
        'view_name',
        'request_method',
        'request_data_display',
        'ip_address',
        'user_agent',
        'created_at'
    ]

    date_hierarchy = 'created_at'
    list_per_page = 50
    actions = ['mark_as_resolved', 'mark_as_unresolved']

    fieldsets = (
        ('خطا', {
            'fields': ('error_type', 'error_message', 'stack_trace_display')
        }),
        ('محتوا', {
            'fields': ('user', 'url', 'view_name')
        }),
        ('درخواست', {
            'fields': ('request_method', 'request_data_display'),
            'classes': ('collapse',)
        }),
        ('سیستم', {
            'fields': ('ip_address', 'user_agent'),
            'classes': ('collapse',)
        }),
        ('وضعیت', {
            'fields': ('is_resolved', 'resolved_at', 'resolved_by')
        }),
        ('زمان', {
            'fields': ('created_at',)
        })
    )

    def error_badge(self, obj):
        """Display error type"""
        colors = {
            'validation': '#f59e0b',
            'database': '#ef4444',
            'payment': '#dc2626',
            'api': '#ec4899',
            'server': '#7c3aed',
            'other': '#6b7280'
        }

        return format_html(
            '<span style="background: {}; color: white; padding: 4px 12px; '
            'border-radius: 8px; font-size: 11px; font-weight: 600;">{}</span>',
            colors.get(obj.error_type, '#6b7280'),
            obj.get_error_type_display()
        )

    error_badge.short_description = 'نوع'

    def error_message_short(self, obj):
        """Truncated error message"""
        if len(obj.error_message) > 80:
            return obj.error_message[:80] + '...'
        return obj.error_message

    error_message_short.short_description = 'پیام خطا'

    def user_display(self, obj):
        """Display user"""
        if obj.user:
            url = f'/admin/accounts/user/{obj.user.id}/change/'
            return format_html('<a href="{}">{}</a>', url, obj.user.full_name)
        return format_html('<span style="color: #9ca3af;">-</span>')

    user_display.short_description = 'کاربر'

    def resolution_status(self, obj):
        """Display resolution status"""
        if obj.is_resolved:
            return format_html(
                '<span style="color: #10b981; font-weight: 600;">✓ برطرف شده</span>'
            )
        return format_html(
            '<span style="color: #ef4444; font-weight: 600;">✗ فعال</span>'
        )

    resolution_status.short_description = 'وضعیت'

    def stack_trace_display(self, obj):
        """Display stack trace"""
        if obj.stack_trace:
            return format_html(
                '<pre style="background: #1f2937; color: #f9fafb; padding: 15px; '
                'border-radius: 8px; max-height: 500px; overflow-y: auto; '
                'font-size: 12px; line-height: 1.5;">{}</pre>',
                obj.stack_trace
            )
        return '-'

    stack_trace_display.short_description = 'Stack Trace'

    def request_data_display(self, obj):
        """Display request data"""
        if obj.request_data:
            try:
                formatted = json.dumps(obj.request_data, indent=2, ensure_ascii=False)
                return format_html('<pre>{}</pre>', formatted)
            except:
                return str(obj.request_data)
        return '-'

    request_data_display.short_description = 'داده‌های درخواست'

    @admin.display(description='تاریخ', ordering='created_at')
    def get_jalali_created_at(self, obj):
        return datetime2jalali(obj.created_at).strftime('%Y/%m/%d - %H:%M:%S')

    @admin.action(description='✓ علامت‌گذاری به عنوان برطرف شده')
    def mark_as_resolved(self, request, queryset):
        updated = queryset.update(
            is_resolved=True,
            resolved_at=timezone.now(),
            resolved_by=request.user
        )
        self.message_user(request, f'{updated} خطا به عنوان برطرف شده علامت‌گذاری شد.')

    @admin.action(description='✗ علامت‌گذاری به عنوان برطرف نشده')
    def mark_as_unresolved(self, request, queryset):
        updated = queryset.update(
            is_resolved=False,
            resolved_at=None,
            resolved_by=None
        )
        self.message_user(request, f'{updated} خطا به عنوان برطرف نشده علامت‌گذاری شد.')

    def has_add_permission(self, request):
        return False
